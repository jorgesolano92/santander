# -*- coding: utf-8 -*-
"""Genera el Manual de Operación e Instalación (Word)."""
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parent / "Manual_Operacion_Instalacion_Control_Accesos.docx"


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_heading_styled(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0xC4, 0x14, 0x25) if level == 1 else RGBColor(0x33, 0x33, 0x33)
    return h


def add_para(doc, text, bold=False, size=11, space_after=8):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(text, style="List Bullet")
    if level:
        p.paragraph_format.left_indent = Cm(1.25 * level)
    for run in p.runs:
        set_run_font(run, size=11)
    return p


def add_code(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    run.font.size = Pt(9)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Cm(0.5)
    return p


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for run in p.runs:
                set_run_font(run, size=10, bold=True)
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            cells[c_idx].text = str(val)
            for p in cells[c_idx].paragraphs:
                for run in p.runs:
                    set_run_font(run, size=9)
    doc.add_paragraph()
    return table


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    # Portada
    for _ in range(3):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("MANUAL DE OPERACIÓN E INSTALACIÓN")
    set_run_font(r, size=26, bold=True, color=RGBColor(0xC4, 0x14, 0x25))

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Sistema de Control de Accesos\nBanco Santander / SAIMA SEGURIDAD")
    set_run_font(r, size=16, bold=True)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(
        f"Sistema local (panel de oficina) · Centro de Control (COCE) · App Android (tablet)\n"
        f"Versión documental: 1.0 — {date.today().strftime('%d/%m/%Y')}"
    )
    set_run_font(r, size=11)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = note.add_run(
        "Documento orientado a instalación, configuración, uso operativo y actualizaciones.\n"
        "Incluye tres bloques independientes: Sistema local, COCE y App Android."
    )
    set_run_font(r, size=10)

    doc.add_page_break()

    # Índice
    add_heading_styled(doc, "Índice", 1)
    for item in [
        "1. Visión general del sistema",
        "2. Arquitectura y comunicaciones",
        "3. Puertos y servicios",
        "4. SISTEMA LOCAL (PC industrial de oficina)",
        "5. CENTRO DE CONTROL (COCE)",
        "6. APP ANDROID (tablet Akuvox)",
        "7. Actualizaciones remotas (COCE → sucursales)",
        "8. Modos operativos y funcionamiento",
        "9. Zaguán, llamadas a tablet e intercom",
        "10. Integración Panphone / CSIP (opcional)",
        "11. Checklist de puesta en marcha",
        "12. Solución de problemas frecuentes",
        "13. Anexos (variables de entorno y rutas)",
    ]:
        add_para(doc, item, size=11, space_after=4)

    doc.add_page_break()

    # 1
    add_heading_styled(doc, "1. Visión general del sistema", 1)
    add_para(
        doc,
        "El sistema de control de accesos gestiona las puertas de zaguán de oficinas del Banco Santander. "
        "Se organiza en tres piezas de software que colaboran entre sí:",
    )
    add_table(
        doc,
        ["Pieza", "Dónde corre", "Función"],
        [
            [
                "Sistema local (panel)",
                "PC industrial de cada oficina",
                "Lógica de modos, Modbus a placas ETD8A12, API tablet, panel web, zaguán ESP32, enlace a COCE",
            ],
            [
                "COCE (API + dashboard)",
                "PC/servidor del Centro de Control",
                "Supervisión de sucursales, mensajes, técnicos, cambio de modo remoto, actualizaciones",
            ],
            [
                "App Android",
                "Tablet Akuvox en oficina",
                "Selección de modo, apertura, vídeo RTSP, intercom, recepción de llamadas y mensajes COCE",
            ],
        ],
    )
    add_para(
        doc,
        "El COCE nunca habla directamente con las tablets ni con las placas Modbus. "
        "Siempre pasa por el backend de la sucursal. Las tablets solo hablan con el PC de oficina.",
    )

    # 2
    add_heading_styled(doc, "2. Arquitectura y comunicaciones", 1)
    add_para(doc, "Flujo lógico resumido:", bold=True)
    add_code(
        doc,
        "Tablet Android  ←→  REST/JWT + WebSocket  ←→  PC oficina (FastAPI :8000)\n"
        "                                              │\n"
        "                     ┌────────────────────────┼────────────────────────┐\n"
        "                     ▼                        ▼                        ▼\n"
        "               ETD8A12 Modbus            ESP32 zaguán            COCE (:9000)\n"
        "               (Central/Calle/Oficina)   (pulsadores/LED)        vía WebSocket\n"
        "                     │\n"
        "               audio_bridge :8765  ←→  cámaras (intercom)\n"
        "               vídeo RTSP :554     ←→  tablet (videoportero)",
    )
    add_para(doc, "Capas:", bold=True)
    add_bullet(doc, "Capa UI: tablet Android + panel web del operador en el PC.")
    add_bullet(doc, "Capa lógica: servicio Python (FastAPI) en el PC industrial.")
    add_bullet(doc, "Capa hardware: 3 módulos ETD8A12 (Modbus TCP o RTU) + nodos ESP32 del zaguán.")
    add_bullet(doc, "Capa central: COCE API + dashboard (supervisión multi-sucursal).")

    # 3
    add_heading_styled(doc, "3. Puertos y servicios", 1)
    add_table(
        doc,
        ["Puerto", "Servicio"],
        [
            ["8000", "Backend FastAPI de sucursal (API + panel web en producción)"],
            ["5173", "Frontend panel en desarrollo (Vite)"],
            ["5174", "Dashboard COCE en desarrollo (Vite)"],
            ["9000", "API COCE central"],
            ["8765", "Puente de audio intercom (audio_bridge.py)"],
            ["502 / 5000", "Modbus TCP de placas ETD8A12 (según firmware/config)"],
            ["COM7 @ 9600", "Modbus RTU por defecto (si MODBUS_MODE=rtu)"],
            ["80", "API HTTP de ESP32 zaguán"],
            ["8266", "OTA TCP del firmware ESP32"],
            ["554", "RTSP de cámaras / videoportero"],
            ["8090", "API CSIP / Panphone (si se usa)"],
        ],
    )
    add_para(
        doc,
        "Para acceso remoto entre COCE y oficinas (y tablets fuera de LAN) se usa típicamente ZeroTier "
        "u otra VPN. La lógica de la aplicación no cambia: solo la topología de red.",
    )

    doc.add_page_break()

    # ============================================================
    # 4 SISTEMA LOCAL
    # ============================================================
    add_heading_styled(doc, "4. SISTEMA LOCAL (PC industrial de oficina)", 1)
    add_para(
        doc,
        "Incluye el backend Python (FastAPI), el frontend del panel de configuración, "
        "la comunicación Modbus, el orquestador de zaguán y el cliente WebSocket hacia COCE.",
    )

    add_heading_styled(doc, "4.1 Obtención / descarga", 2)
    add_bullet(doc, "Repositorio: https://github.com/jorgesolano92/santander.git")
    add_bullet(doc, "Clonar en el PC industrial: git clone <url> y checkout de la rama de despliegue acordada.")
    add_bullet(
        doc,
        "Actualizaciones posteriores de panel: preferiblemente desde COCE (zip remoto), sin git pull en oficina "
        "(ver sección 7).",
    )

    add_heading_styled(doc, "4.2 Requisitos", 2)
    add_bullet(doc, "Windows (PC industrial / Windows IoT según contrato).")
    add_bullet(doc, "Python 3.9+ (recomendado 3.11/3.12).")
    add_bullet(doc, "Node.js 18+ (solo para construir el frontend).")
    add_bullet(doc, "Red Ethernet local con IPs estáticas para PC, placas y cámaras.")
    add_bullet(doc, "Adaptador RS-485 o acceso TCP a las 3 placas ETD8A12.")

    add_heading_styled(doc, "4.3 Instalación (desarrollo)", 2)
    add_para(doc, "Backend:", bold=True)
    add_code(
        doc,
        "cd backend\n"
        "python -m venv venv\n"
        "venv\\Scripts\\activate\n"
        "pip install -r requirements.txt\n"
        "copy .env.example .env\n"
        "python run.py",
    )
    add_para(doc, "URL: http://localhost:8000 — Documentación OpenAPI: http://localhost:8000/docs")

    add_para(doc, "Frontend (panel web):", bold=True)
    add_code(
        doc,
        "cd frontend\n"
        "npm install\n"
        "npm run dev",
    )
    add_para(doc, "URL: http://localhost:5173 (Vite hace proxy de /api al puerto 8000).")

    add_heading_styled(doc, "4.4 Instalación / arranque en producción", 2)
    add_para(
        doc,
        "Se construye el frontend y el backend sirve la SPA y la API en el mismo origen (:8000):",
    )
    add_code(
        doc,
        "cd frontend\n"
        "npm install\n"
        "npm run build\n"
        "\n"
        "cd ..\\backend\n"
        "venv\\Scripts\\activate\n"
        "python run.py\n"
        "# equivalente: uvicorn app.main:app --host 0.0.0.0 --port 8000",
    )
    add_para(
        doc,
        "Opcional: copiar frontend/dist a backend/static y definir STATIC_DIR=./static en .env. "
        "En producción el proceso debería ejecutarse como servicio Windows (NSSM, WinSW o similar) "
        "para arranque automático; el repositorio documenta el reinicio del servicio tras updates, "
        "aunque el instalador MSI/servicio no forma parte del paquete de código actual.",
    )

    add_heading_styled(doc, "4.5 Configuración esencial (.env)", 2)
    add_para(doc, "Archivo: backend/.env (partir de backend/.env.example). Principales variables:")
    add_table(
        doc,
        ["Variable", "Uso"],
        [
            ["DATABASE_URL", "SQLite local (p.ej. sqlite:///./data/control_accesos.db)"],
            ["MODBUS_MODE", "tcp o rtu"],
            ["MODBUS_SERIAL_*", "Puerto COM, baudios, etc. (si RTU)"],
            ["TABLET_JWT_SECRET / PANEL_JWT_SECRET", "Secretos JWT (obligatorios en producción)"],
            ["TABLET_JWT_EXPIRE_MINUTES", "Caducidad token tablet (p.ej. 10080 = 7 días)"],
            ["COCE_WS_ENABLED / COCE_WS_URL", "Enlace saliente al Centro de Control"],
            ["COCE_INSTALLATION_ID / COCE_INGEST_TOKEN", "Identidad de la sucursal en COCE"],
            ["COCE_MESSAGE_SEND_WEB / TABLET", "Reparto de mensajes a panel y/o tablets"],
            ["ZAGUAN_*", "IP ESP32, captura-only, etc."],
            ["CSIP_*", "Integración Panphone (opcional)"],
            ["TABLET_CALL_*", "Llamadas entrantes a tablets"],
        ],
    )
    add_para(
        doc,
        "La configuración de placas, reglas, horarios y zaguán se gestiona también desde el panel web "
        "y se persiste en SQLite (backend/data/control_accesos.db). "
        "Overrides de ESP32 por canal: backend/data/zaguan_device_target.json.",
    )

    add_heading_styled(doc, "4.6 Primer acceso al panel web", 2)
    add_bullet(doc, "Abrir http://<IP-PC>:8000 (producción) o :5173 (desarrollo).")
    add_bullet(doc, "Registrar el primer usuario de panel (flujo de setup / PANEL_SETUP_TOKEN según .env).")
    add_bullet(doc, "Configurar las 3 placas ETD8A12 (IP/puerto o slave_id), reglas y modos.")
    add_bullet(doc, "Verificar lectura de entradas/salidas en el panel.")

    add_heading_styled(doc, "4.7 Actualización del sistema local", 2)
    add_para(doc, "Opción A — Manual en el PC:", bold=True)
    add_bullet(doc, "git pull (o copiar release).")
    add_bullet(doc, "Actualizar dependencias Python/npm si cambian.")
    add_bullet(doc, "npm run build en frontend.")
    add_bullet(doc, "Reiniciar el servicio / proceso Python.")

    add_para(doc, "Opción B — Remota desde COCE (recomendada en operación):", bold=True)
    add_bullet(doc, "COCE publica release kind=panel.")
    add_bullet(doc, "En el panel de sucursal aparece banner → Actualizar panel.")
    add_bullet(doc, "Se descarga zip, se verifica SHA-256 y se aplica sin tocar data/, BD ni .env.")
    add_bullet(doc, "Reiniciar servicio Windows y recargar el navegador.")

    add_heading_styled(doc, "4.8 Componentes internos del sistema local", 2)
    add_table(
        doc,
        ["Componente", "Ruta / nota"],
        [
            ["API FastAPI", "backend/app/main.py"],
            ["Rutas tablet JWT", "/api/v1/... y WS /api/v1/ws/calls"],
            ["Rutas panel", "/api/panel/... y WS /api/panel/ws/live"],
            ["Orquestador zaguán", "backend/app/services/zaguan_orchestrator.py"],
            ["Hub llamadas tablet", "backend/app/services/tablet_call_hub.py"],
            ["Hardware Modbus", "backend/app/hardware/"],
            ["Cliente COCE", "backend/app/coce/"],
            ["Módulo CSIP Panphone", "backend/app/csip/"],
            ["Firmware / OTA ESP32", "backend/zaguan_*.ino, python_ota_update.py"],
        ],
    )

    doc.add_page_break()

    # ============================================================
    # 5 COCE
    # ============================================================
    add_heading_styled(doc, "5. CENTRO DE CONTROL (COCE)", 1)
    add_para(
        doc,
        "El COCE es el punto central de supervisión. Consta de coce-api (FastAPI + SQLite) y "
        "coce-dashboard (React). Las sucursales mantienen un WebSocket saliente hacia el COCE.",
    )

    add_heading_styled(doc, "5.1 Obtención / descarga", 2)
    add_bullet(doc, "Mismo repositorio santander: carpetas coce-api/ y coce-dashboard/.")
    add_bullet(doc, "Se despliega en la máquina del Centro de Control (IP conocida por las oficinas).")

    add_heading_styled(doc, "5.2 Requisitos", 2)
    add_bullet(doc, "Python 3.9+ y Node.js 18+.")
    add_bullet(doc, "Red alcanzable por las sucursales (ZeroTier/VPN o IP pública con reglas de firewall).")
    add_bullet(doc, "Secretos fuertes: COCE_JWT_SECRET y COCE_SECRETS_KEY (Fernet).")

    add_heading_styled(doc, "5.3 Instalación — API", 2)
    add_code(
        doc,
        "cd coce-api\n"
        "python -m venv .venv\n"
        ".venv\\Scripts\\activate\n"
        "pip install -r requirements.txt\n"
        "copy .env.example .env\n"
        "# Editar: COCE_JWT_SECRET, COCE_SECRETS_KEY, COCE_CORS_ORIGINS\n"
        "python -m app.main",
    )
    add_para(doc, "Escucha por defecto en http://0.0.0.0:9000 — Health: http://<host>:9000/health")

    add_para(doc, "Generar clave Fernet:", bold=True)
    add_code(
        doc,
        'python -c "from cryptography.fernet import Fernet; print(Fernet.generate_key().decode())"',
    )

    add_heading_styled(doc, "5.4 Instalación — Dashboard", 2)
    add_code(
        doc,
        "cd coce-dashboard\n"
        "npm install\n"
        "copy .env.example .env\n"
        "# VITE_COCE_API_URL=http://localhost:9000\n"
        "npm run dev",
    )
    add_para(doc, "Desarrollo: http://localhost:5174")
    add_para(doc, "Producción:")
    add_code(doc, "npm run build\n# servir dist/ con el mismo origen o CORS permitido")

    add_heading_styled(doc, "5.5 Primer usuario y sucursales", 2)
    add_bullet(doc, "Con BD vacía: registrar usuario vía UI o POST /api/coce/auth/register.")
    add_bullet(doc, "Si ya hay usuarios: cabecera X-Coce-Setup-Token = COCE_SETUP_TOKEN.")
    add_bullet(doc, "Login → JWT de sesión (dashboard en sessionStorage).")
    add_bullet(
        doc,
        "Crear sucursales: se genera installation_id + ingest_token (mostrar una sola vez). "
        "Esos valores se copian al .env del backend de la oficina (COCE_INSTALLATION_ID, COCE_INGEST_TOKEN, COCE_WS_URL).",
    )

    add_heading_styled(doc, "5.6 Funcionalidades operativas", 2)
    add_table(
        doc,
        ["Función", "Descripción"],
        [
            ["Mapa / listado sucursales", "Estado online/offline por heartbeat"],
            ["Snapshot remoto", "Modo activo y estado de placas vía proxy HTTP al panel"],
            ["Cambio de modo remoto", "POST set-mode con auditoría"],
            ["Mensajería", "Mensajes a sucursal → panel y/o tablets"],
            ["Técnicos", "Sincronización de técnicos hacia oficinas"],
            ["Alertas", "Eventos / avisos centralizados"],
            ["Actualizaciones", "Publicar panel zip o APK y desplegar por WS"],
            ["Auditoría", "Histórico de acciones administrativas (no editable desde sucursal)"],
        ],
    )

    add_heading_styled(doc, "5.7 Actualización del propio COCE", 2)
    add_bullet(doc, "En la máquina COCE: git pull del repositorio.")
    add_bullet(doc, "Actualizar dependencias; reiniciar coce-api.")
    add_bullet(doc, "Rebuild del dashboard (npm run build) y redesplegar estáticos.")
    add_bullet(doc, "No hay auto-update del COCE desde fuera; es el origen de los paquetes hacia sucursales.")

    add_heading_styled(doc, "5.8 Enlace sucursal ↔ COCE", 2)
    add_para(doc, "En cada oficina (backend/.env):")
    add_code(
        doc,
        "COCE_WS_ENABLED=true\n"
        "COCE_WS_URL=ws://<IP-COCE>:9000/api/coce/ws/branch/<installation_id>\n"
        "COCE_INSTALLATION_ID=<uuid>\n"
        "COCE_INGEST_TOKEN=<token>\n"
        "COCE_HEARTBEAT_INTERVAL_SECONDS=60\n"
        "COCE_MESSAGE_SEND_TABLET=true\n"
        "COCE_MESSAGE_SEND_WEB=true",
    )
    add_para(
        doc,
        "Tipos de mensajes por el canal WS (ejemplos): heartbeat, modos, coce_message, "
        "technicians_sync, software_update, update_status.",
    )

    doc.add_page_break()

    # ============================================================
    # 6 APP ANDROID
    # ============================================================
    add_heading_styled(doc, "6. APP ANDROID (tablet Akuvox)", 1)
    add_para(
        doc,
        "Aplicación Expo / React Native con módulos nativos Android (vídeo/SDK). "
        "Repositorio separado: https://github.com/jorgesolano92/santander-app.git",
    )

    add_heading_styled(doc, "6.1 Obtención / descarga", 2)
    add_bullet(doc, "Código fuente: clonar el repositorio santander-app.")
    add_bullet(doc, "Instalación en tablets: APK firmada (debug o release) vía sideload / software Akuvox.")
    add_bullet(doc, "Distribución operativa: COCE sube la APK → panel de oficina descarga → operador instala en tablets.")

    add_heading_styled(doc, "6.2 Requisitos de desarrollo", 2)
    add_bullet(doc, "Node.js 18+, npm.")
    add_bullet(doc, "Android Studio / SDK / JDK para builds nativos.")
    add_bullet(doc, "Dispositivo o emulador; para intercom nativo preferible tablet real.")
    add_bullet(doc, "Opcional: EAS CLI para builds en la nube (eas.json).")

    add_heading_styled(doc, "6.3 Instalación en entorno de desarrollo", 2)
    add_code(
        doc,
        "cd santander-app\n"
        "npm install\n"
        "npm run dev                 # Expo\n"
        "npm run android             # expo run:android (nativo)\n"
        "npm run android:tablet      # build/install orientado a tablet",
    )

    add_heading_styled(doc, "6.4 Compilación de APK", 2)
    add_para(doc, "Opción recomendada (Gradle, conserva código nativo SDK):", bold=True)
    add_code(
        doc,
        "cd android\n"
        ".\\gradlew clean\n"
        ".\\gradlew assembleDebug\n"
        "# o release:\n"
        ".\\gradlew assembleRelease\n"
        "\n"
        "# APK debug:\n"
        "# android\\app\\build\\outputs\\apk\\debug\\app-debug.apk\n"
        "# Instalar:\n"
        ".\\gradlew installDebug",
    )
    add_para(doc, "Con Expo:")
    add_code(doc, "npx expo run:android")
    add_para(doc, "Con EAS (perfiles en eas.json: apk-release, standalone, etc.):")
    add_code(doc, "eas build --platform android --profile apk-release")
    add_para(
        doc,
        "Importante: no usar comandos destructivos (git reset --hard) antes de compilar si hay cambios "
        "locales en el SDK nativo (ver docs/COMPILACION_ANDROID.md).",
    )

    add_heading_styled(doc, "6.5 Configuración de la app", 2)
    add_para(
        doc,
        "La configuración operativa vive en la tablet (AsyncStorage), con defaults en "
        "config/defaultDoorAppConfig.ts. Campos clave:",
    )
    add_bullet(doc, "network.consoleIP — IP del PC industrial (ej. 192.168.1.155).")
    add_bullet(doc, "api.port — 8000.")
    add_bullet(doc, "Puertas: cámara IP, RTSP, bridgeUrl (ws://<PC>:8765), mapeo Modbus.")
    add_bullet(doc, "Modos habilitados/deshabilitados por oficina.")
    add_bullet(doc, "officeWithATM — muestra modo carga cajero si aplica.")
    add_para(
        doc,
        "Intercom actual en operación: modo puente PC (INTERCOM_BRIDGE_ONLY = true en "
        "config/intercomFeatures.ts). SIP/CSIP están preparados en código pero ocultos en UI hasta activarlos.",
    )

    add_heading_styled(doc, "6.6 Puente de audio (PC de oficina)", 2)
    add_para(
        doc,
        "Para intercom bidireccional con cámaras TVT/Safire se ejecuta en el PC (aparte del FastAPI):",
    )
    add_code(
        doc,
        "cd Release_vcx_x64\n"
        "pip install -r bridge_requirements.txt\n"
        "python audio_bridge.py",
    )
    add_bullet(doc, "WebSocket en :8765.")
    add_bullet(doc, "Requiere DLLs del SDK + Visual C++ Redistributable x64.")
    add_bullet(doc, "Ajustes: BRIDGE_ENV.md y bridge_cameras.json.")

    add_heading_styled(doc, "6.7 Autorización de dispositivo", 2)
    add_para(
        doc,
        "La app obtiene el Android ID de la tablet y consulta al backend si el dispositivo está autorizado. "
        "Si no lo está, muestra pantalla de dispositivo no autorizado. "
        "Las tablets deben estar dadas de alta / autorizadas en el panel o flujo configurado.",
    )

    add_heading_styled(doc, "6.8 Funciones principales de la app", 2)
    add_bullet(doc, "Visualización del modo activo y descripción + esquema/diagrama.")
    add_bullet(doc, "Cambio de modo (lista filtrada según modos habilitados en ajustes).")
    add_bullet(doc, "Modo carga cajero: layout con vídeo a la derecha y descripción/esquema a la izquierda.")
    add_bullet(doc, "Modo manual / bloqueo de puertas: control de aperturas e intercom.")
    add_bullet(doc, "Llamadas entrantes (videoportero): suenan todas las tablets; la primera que contesta gana.")
    add_bullet(doc, "Mensajes COCE (toast + historial).")
    add_bullet(doc, "Emergencia e incendio (confirmación y desactivación).")

    add_heading_styled(doc, "6.9 Actualización de la app", 2)
    add_para(doc, "Flujo operativo recomendado:", bold=True)
    add_bullet(doc, "Compilar y firmar APK en el entorno de desarrollo.")
    add_bullet(doc, "En COCE → Actualizaciones: subir APK + versión (ej. 2.8.1) + notas.")
    add_bullet(doc, "Lanzar despliegue a sucursales (kind=tablet_apk).")
    add_bullet(doc, "En el panel de oficina: banner → Descargar APK.")
    add_bullet(doc, "Instalar en cada tablet con el software Akuvox (no hay OTA silenciosa ni Play Store).")
    add_para(
        doc,
        "La tablet no descarga directamente del COCE. La versión la introduce el operador al publicar; "
        "no se lee automáticamente el versionCode del APK.",
    )

    doc.add_page_break()

    # 7 Updates
    add_heading_styled(doc, "7. Actualizaciones remotas (COCE → sucursales)", 1)
    add_para(
        doc,
        "Objetivo: distribuir software sin que las PCs de oficina tengan acceso a GitHub ni hagan git pull.",
    )
    add_table(
        doc,
        ["Tipo", "kind", "Qué actualiza", "Acción en sucursal"],
        [
            ["Panel PC", "panel", "Backend + frontend", "Banner → Actualizar panel → reiniciar servicio"],
            ["APK tablet", "tablet_apk", "APK Android", "Banner → Descargar APK → instalar con Akuvox"],
        ],
    )

    add_heading_styled(doc, "7.1 Publicar panel desde COCE", 2)
    add_bullet(doc, "En COCE: git pull del checkout del panel; opcional npm run build en frontend.")
    add_bullet(doc, "En coce-api/.env: COCE_PANEL_SOURCE_DIR=<ruta al repo santander>.")
    add_bullet(doc, "Dashboard → Actualizaciones → Publicar desde local (empaqueta backend + frontend/dist).")
    add_bullet(doc, "Excluye: data/, *.db, .env, node_modules, .git, venv.")
    add_bullet(doc, "Seleccionar release → Lanzar despliegue (todas o sucursales concretas).")

    add_heading_styled(doc, "7.2 Publicar APK", 2)
    add_bullet(doc, "Subir .apk + versión + changelog.")
    add_bullet(doc, "Artefacto en coce-api/data/releases/ (fuera de git).")
    add_bullet(doc, "Desplegar; las oficinas reciben aviso por WebSocket.")

    add_heading_styled(doc, "7.3 Estados reportados", 2)
    add_para(
        doc,
        "La sucursal informa update_status: available | downloading | applying | success | failed | offline. "
        "El dashboard muestra la cola de despliegues.",
    )

    # 8 Modes
    add_heading_styled(doc, "8. Modos operativos y funcionamiento", 1)
    add_para(
        doc,
        "Existen 7 modos con exclusión mutua (solo uno activo). Se activan por horario/entradas del "
        "módulo Central, por tablet, por panel o por COCE (remoto).",
    )
    add_table(
        doc,
        ["Modo", "Idea operativa"],
        [
            ["Comercial automático", "Flujo normal de atención en horario comercial"],
            ["Comercial esclusa", "Paso controlado tipo esclusa"],
            ["Horario extendido", "Fuera de horario comercial; llamadas tablet en ciertos pulsadores"],
            ["Horario autoservicio", "Autoservicio / acceso restringido según reglas"],
            ["Oficina cerrada", "Cierres de seguridad; puertas cerradas"],
            ["Carga de cajero", "Carga ATM en zaguán; P1 cerrada con llamada; vídeo en home"],
            ["Bloqueo de puertas (manual)", "Control manual desde tablet; llamadas a tablets"],
        ],
    )
    add_para(
        doc,
        "Además: emergencia e incendio (prioridad alta). El detalle de las 33 actuaciones "
        "(ATACA / ACTIVA / DESACTIVA / NO ACTUA SI) está en EXPLICACION_ACTUACIONES_Y_MODOS.md "
        "y en los Excel de actuaciones del proyecto.",
    )

    # 9 Zaguan / calls
    add_heading_styled(doc, "9. Zaguán, llamadas a tablet e intercom", 1)

    add_heading_styled(doc, "9.1 Pulsadores ESP32", 2)
    add_para(doc, "El visitante pulsa un botón en el zaguán:")
    add_code(
        doc,
        "ESP32 → POST /zaguan/pulsacion/p1|p2|p3|p4\n"
        "      → zaguan_orchestrator (según modo)\n"
        "      → abre puerta por Modbus (si aplica) y/o inicia llamada tablet\n"
        "      → actualiza LED del ESP32 (libre/ocupado/abriendo)",
    )
    add_bullet(doc, "p1 / p3 ≈ puerta calle (P1).")
    add_bullet(doc, "p2 / p4 ≈ puerta oficina (P2).")
    add_bullet(doc, "ZAGUAN_PULSACION_CAPTURE_ONLY=1: solo captura, no abre Modbus (banco de pruebas).")

    add_heading_styled(doc, "9.2 Llamadas a tablets (first-wins)", 2)
    add_para(
        doc,
        "En modos configurados (p.ej. manual, carga cajero, extendido), la pulsación dispara "
        "tablet_call_hub.start_call():",
    )
    add_bullet(doc, "Se envía incoming_call por WebSocket a TODAS las tablets conectadas.")
    add_bullet(doc, "Todas suenan / muestran overlay.")
    add_bullet(doc, "La primera que envía answer recibe call_accepted.")
    add_bullet(doc, "Las demás reciben call_taken y dejan de sonar.")
    add_bullet(doc, "Timeout configurable (TABLET_CALL_TIMEOUT).")
    add_para(
        doc,
        "Tras contestar, la tablet abre ManualModeModal y arranca el intercom del modo configurado "
        "(puente por defecto). El canal de intercom bidireccional también es exclusivo "
        "(claim/release): solo una tablet a la vez.",
    )

    add_heading_styled(doc, "9.3 Intercom", 2)
    add_table(
        doc,
        ["Modo", "Estado"],
        [
            ["bridge (puente PC)", "Operación actual: tablet ↔ WS :8765 ↔ audio_bridge ↔ cámara"],
            ["sdk", "Nativo Android; disponible en código, UI condicionada"],
            ["sip / CSIP", "Preparados para Panphone + PBX; UI condicionada (INTERCOM_BRIDGE_ONLY)"],
        ],
    )

    # 10 CSIP
    add_heading_styled(doc, "10. Integración Panphone / CSIP (opcional)", 1)
    add_para(
        doc,
        "Módulo backend aparte (backend/app/csip/) expuesto bajo /api/csip. Pensado para pruebas "
        "y futura operación con placa Panphone + centralita PBX.",
    )
    add_bullet(doc, "Webhooks: POST /api/csip/notify, /notify/p1, /notify/p2 (notification_url en la placa).")
    add_bullet(doc, "Proxy saliente: call_start, led_control, button_event.")
    add_bullet(doc, "CSIP_FORWARD_PULSACION_TO_ZAGUAN=true para reutilizar la lógica de zaguán/llamadas tablet.")
    add_para(
        doc,
        "Patrón recomendado con PBX: la placa notifica al backend; suenan todas las tablets por WebSocket; "
        "solo la que contesta abre la sesión SIP / call_start hacia su extensión. "
        "Detalle técnico en docs/CSIP_SIP_INTEGRATION.md de la app.",
    )

    doc.add_page_break()

    # 11 Checklist
    add_heading_styled(doc, "11. Checklist de puesta en marcha", 1)
    add_para(doc, "Orden sugerido:", bold=True)
    add_bullet(doc, "1. Red local / ZeroTier: IPs de PC, placas, cámaras, ESP32, tablets.")
    add_bullet(doc, "2. Instalar y arrancar sistema local (backend + frontend build) en :8000.")
    add_bullet(doc, "3. Configurar Modbus (TCP o RTU) y verificar I/O en panel.")
    add_bullet(doc, "4. Flashear/configurar ESP32 zaguán y apuntar al backend; probar pulsaciones.")
    add_bullet(doc, "5. Arrancar audio_bridge :8765 si se usa intercom puente.")
    add_bullet(doc, "6. Instalar COCE (api :9000 + dashboard); crear usuario y sucursal.")
    add_bullet(doc, "7. Cablear COCE_WS_* en la oficina y verificar heartbeat online.")
    add_bullet(doc, "8. Compilar e instalar APK; apuntar consoleIP al PC; autorizar dispositivo.")
    add_bullet(doc, "9. Probar cambio de modo, apertura, llamada tablet, mensaje COCE.")
    add_bullet(doc, "10. Probar ciclo de actualización panel y APK desde COCE (banco).")

    # 12 Troubleshooting
    add_heading_styled(doc, "12. Solución de problemas frecuentes", 1)
    add_table(
        doc,
        ["Síntoma", "Comprobación"],
        [
            ["Tablet WS 403 / se desconecta", "JWT caducado o inválido; revisar TABLET_JWT_* y re-login"],
            ["No cambia modo / modo deshabilitado", "Ajustes de tablet: modes.*.enabled; el modo no debe listarse si está off"],
            ["Placa Modbus desconectada", "Cableado RS-485/TCP, COM, slave_id, timeouts"],
            ["COCE offline", "COCE_WS_URL alcanzable, token e installation_id correctos"],
            ["No suena llamada en tablets", "Modo actual en TABLET_CALL_MODES; WS conectado; pulsador permitido"],
            ["Sin audio intercom", "audio_bridge en marcha :8765; bridgeUrl; DLLs/VC++"],
            ["CSIP notify abre el panel en navegador", "Normal en GET; el webhook es POST"],
            ["Crash Android al abrir modos", "SVG orient=auto-start-reverse; ya mitigado en ModeIcon"],
            ["Update panel no aplica", "Banner apply; SHA; espacio disco; reinicio servicio tras apply"],
        ],
    )

    # 13 Anexos
    add_heading_styled(doc, "13. Anexos", 1)

    add_heading_styled(doc, "13.1 Documentación de referencia en los repositorios", 2)
    add_table(
        doc,
        ["Documento", "Contenido"],
        [
            ["santander/README.md", "Arranque rápido"],
            ["santander/DESARROLLO.md", "Dev y build producción"],
            ["santander/ARQUITECTURA.md", "Capas y persistencia"],
            ["santander/ACTUALIZACIONES_REMOTAS.md", "Contrato updates COCE"],
            ["santander/EXPLICACION_ACTUACIONES_Y_MODOS.md", "Modos y actuaciones"],
            ["santander/coce-api/README.md", "API central"],
            ["santander/coce-dashboard/README.md", "UI COCE"],
            ["santander/backend/SAIMA_*.md", "ESP32 instalación/API"],
            ["santander-app/docs/COMPILACION_ANDROID.md", "Build APK"],
            ["santander-app/docs/CSIP_SIP_INTEGRATION.md", "SIP / CSIP"],
            ["santander-app/Release_vcx_x64/BRIDGE_ENV.md", "Puente audio"],
        ],
    )

    add_heading_styled(doc, "13.2 Rutas útiles", 2)
    add_code(
        doc,
        "Panel oficina:     http://<IP-PC>:8000\n"
        "OpenAPI oficina:   http://<IP-PC>:8000/docs\n"
        "COCE API:          http://<IP-COCE>:9000\n"
        "COCE dashboard:    http://<IP-COCE>:5174 (dev) o URL de producción\n"
        "Zaguan ESP:        http://<IP-ESP>/ ...  OTA TCP :8266\n"
        "Audio bridge:      ws://<IP-PC>:8765",
    )

    add_heading_styled(doc, "13.3 Límites actuales a tener en cuenta", 2)
    add_bullet(doc, "No hay Docker/compose ni MSI empaquetado en el repositorio.")
    add_bullet(doc, "El servicio Windows se asume en operación pero el instalador no viene en el código.")
    add_bullet(doc, "La APK no se autoinstala en Akuvox; requiere intervención del operador.")
    add_bullet(doc, "SIP/Panphone: integración preparada; operación diaria actual usa puente PC.")

    add_para(doc, "")
    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run(
        "— Fin del documento —\n"
        "Control de Accesos Santander / SAIMA · Manual de operación e instalación v1.0"
    )
    set_run_font(r, size=10, color=RGBColor(0x66, 0x66, 0x66))

    doc.save(OUT)
    print(f"OK: {OUT}")
    return OUT


if __name__ == "__main__":
    build()
